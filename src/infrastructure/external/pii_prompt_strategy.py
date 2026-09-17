"""Strategy for loading and rendering prompts from Google Sheets with local Excel fallback."""

import logging
import os
from pathlib import Path
from typing import Dict, Any, List, Optional, Tuple
from jinja2 import Environment
from config.config import get_config

logger = logging.getLogger(__name__)

DEFAULT_EXCEL_PATH = 'src/prompts/prompts_dev.xlsx'

WORKSHEET_ALIASES: Dict[str, List[str]] = {
    'personal_data_detection': ['personal_data_detection', 'pii detection', 'pii_detection', 'detection'],
    'personal_data_reflection': ['personal_data_reflection', 'pii reflection', 'pii_reflection', 'reflection'],
    'non_personal_data_classificatio': [
        'non_personal_data_classificatio',
        'non_personal_data_classification',
        'non-pii classification',
        'non_pii_classification',
        'non_pii',
    ],
    'non_personal_data_default_class': [
        'non_personal_data_default_class',
        'non_personal_data_default_classification',
        'non-pii default',
        'non_pii_default',
        'default',
    ],
    'readme': ['readme', 'readme scan', 'readme_scan'],
}


def _parse_section_id(val: Any) -> Tuple[int, Any]:
    if val is None:
        return (2, '')
    s = str(val).strip()
    try:
        return (0, int(s))
    except ValueError:
        try:
            return (0, float(s))
        except ValueError:
            return (1, s)


class SpreadsheetPromptStrategy:
    """
    Unified strategy to retrieve prompt rules from Google Sheets with local Excel fallback.

    Reads a worksheet matching `worksheet_name` or its aliases, filters out rows where `enabled` is False,
    sorts strictly by `section_id`, and extracts content rules into template strings or dictionaries.
    """

    def __init__(
        self,
        worksheet_name: str,
        spreadsheet_url: Optional[str] = None,
        excel_path: Optional[str] = None,
        store: Optional[Any] = None,
        cache_key: Optional[str] = None,
    ):
        self.worksheet_name = worksheet_name
        self.spreadsheet_url = spreadsheet_url or get_config().GOOGLE_SHEET_URL
        self.excel_path = excel_path or os.getenv('PROMPTS_EXCEL_PATH', DEFAULT_EXCEL_PATH)
        self.store = store
        self.cache_key = cache_key or f'prompt_cache_{worksheet_name}'
        self.rules_cache_key = f'{self.cache_key}_rules'
        self._cached_template_str: Optional[str] = None
        self._cached_rules: Optional[List[Dict[str, Any]]] = None
        self._jinja_env = Environment(trim_blocks=True, lstrip_blocks=True)

    def _matches_worksheet_name(self, title: str) -> bool:
        t_clean = title.strip().lower()
        ws_clean = self.worksheet_name.strip().lower()
        if t_clean == ws_clean or t_clean.startswith(ws_clean):
            return True
        aliases = WORKSHEET_ALIASES.get(ws_clean, [ws_clean])
        for alias in aliases:
            if t_clean == alias or t_clean.startswith(alias):
                return True
        return False

    def _fetch_from_google_sheets(self) -> Optional[List[List[Any]]]:
        from src.infrastructure.external.google_sheets_client import get_gsheets

        try:
            spreadsheet = get_gsheets().open_by_url(self.spreadsheet_url)
            # Try exact worksheet lookup first
            try:
                worksheet = spreadsheet.worksheet(self.worksheet_name)
                return worksheet.get_all_values()
            except Exception:
                # Alias / fuzzy lookup fallback
                matching_ws = [w for w in spreadsheet.worksheets() if self._matches_worksheet_name(w.title)]
                if matching_ws:
                    return matching_ws[0].get_all_values()
                raise
        except Exception as e:
            logger.warning(f'Failed to fetch prompt rules from Google Sheet ({self.worksheet_name}): {e}')
            return None

    def _fetch_from_local_excel(self) -> Optional[List[List[Any]]]:
        excel_file = Path(self.excel_path)
        if not excel_file.exists():
            logger.warning(f'Local Excel fallback file not found at: {self.excel_path}')
            return None
        try:
            import openpyxl

            wb = openpyxl.load_workbook(str(excel_file), data_only=True)
            target_sheet = None
            for sname in wb.sheetnames:
                if self._matches_worksheet_name(sname):
                    target_sheet = wb[sname]
                    break
            if target_sheet is None:
                logger.warning(f'No matching sheet in Excel file "{self.excel_path}" for "{self.worksheet_name}"')
                return None

            values = list(target_sheet.iter_rows(values_only=True))
            return values
        except Exception as e:
            logger.error(f'Failed to read local Excel fallback file ({self.excel_path}): {e}')
            return None

    def load_rows(self, force_refresh: bool = False) -> Optional[List[List[Any]]]:
        """Fetch raw rows: first try Google Sheets, then fallback to local Excel."""
        values = self._fetch_from_google_sheets()
        if values:
            logger.info(f'Loaded prompt rules for "{self.worksheet_name}" from Google Sheets.')
            return values

        logger.info(f'Attempting local Excel fallback for "{self.worksheet_name}".')
        values = self._fetch_from_local_excel()
        if values:
            logger.info(f'Loaded prompt rules for "{self.worksheet_name}" from local Excel ({self.excel_path}).')
            return values

        return None

    def load_rules(self, force_refresh: bool = False) -> List[Dict[str, Any]]:
        """Parse, filter by enabled, and sort strictly by section_id."""
        if self._cached_rules is not None and not force_refresh:
            return self._cached_rules

        if self.store and not force_refresh:
            try:
                get_fn = getattr(self.store, 'get_object', getattr(self.store, 'get', None))
                cached_rules = get_fn(self.rules_cache_key) if get_fn else None
                if cached_rules:
                    logger.info(f'Loaded prompt rules for "{self.worksheet_name}" from Redis cache.')
                    self._cached_rules = cached_rules
                    return self._cached_rules
            except Exception as e:
                logger.error(f'Failed to load prompt rules from Redis cache: {e}')

        values = self.load_rows(force_refresh=force_refresh)
        if not values or len(values) < 2:
            return []

        header = [str(h).strip().lower() if h is not None else '' for h in values[0]]
        rows = values[1:]

        if 'content' not in header:
            logger.error(f'Worksheet "{self.worksheet_name}" is missing required column "content"')
            return []

        content_idx = header.index('content')
        enabled_idx = header.index('enabled') if 'enabled' in header else None
        section_id_idx = header.index('section_id') if 'section_id' in header else None
        category_idx = header.index('category') if 'category' in header else None
        qa_notes_idx = header.index('qa_notes') if 'qa_notes' in header else None

        # Filter enabled
        if enabled_idx is not None:
            rows = [
                r
                for r in rows
                if enabled_idx < len(r) and r[enabled_idx] not in (None, False, 'False', 'false', 0, '0', 'FALSE')
            ]

        # Sort strictly by section_id
        if section_id_idx is not None:
            rows = sorted(
                rows, key=lambda r: _parse_section_id(r[section_id_idx]) if section_id_idx < len(r) else (2, '')
            )

        parsed_rules = []
        for r in rows:
            if content_idx < len(r) and r[content_idx]:
                rule_text = str(r[content_idx]).strip().replace('\\n', '\n')
                if rule_text:
                    rule_obj = {'content': rule_text}
                    if section_id_idx is not None and section_id_idx < len(r):
                        rule_obj['section_id'] = r[section_id_idx]
                    if category_idx is not None and category_idx < len(r):
                        rule_obj['category'] = r[category_idx]
                    if qa_notes_idx is not None and qa_notes_idx < len(r):
                        rule_obj['qa_notes'] = r[qa_notes_idx]
                    parsed_rules.append(rule_obj)

        self._cached_rules = parsed_rules

        if self.store and self.rules_cache_key and parsed_rules:
            try:
                set_fn = getattr(self.store, 'set_object', getattr(self.store, 'set', None))
                if set_fn:
                    set_fn(self.rules_cache_key, parsed_rules, expire_in_seconds=60 * 60 * 12)
            except Exception as e:
                logger.error(f'Failed to set prompt rules in Redis cache: {e}')

        return self._cached_rules

    def load_template_string(self, force_refresh: bool = False) -> Optional[str]:
        if self._cached_template_str is not None and not force_refresh:
            return self._cached_template_str

        if self.store and not force_refresh:
            try:
                get_fn = getattr(self.store, 'get_object', getattr(self.store, 'get', None))
                cached_val = get_fn(self.cache_key) if get_fn else None
                if cached_val:
                    self._cached_template_str = cached_val
                    return self._cached_template_str
            except Exception as e:
                logger.error(f'Failed to load prompt template from Redis cache: {e}')

        rules = self.load_rules(force_refresh=force_refresh)
        if not rules:
            return None

        prompt_parts = [r['content'] for r in rules]
        template_str = '\n\n'.join(prompt_parts)
        self._cached_template_str = template_str

        if self.store and self.cache_key:
            try:
                set_fn = getattr(self.store, 'set_object', getattr(self.store, 'set', None))
                if set_fn:
                    set_fn(self.cache_key, template_str, expire_in_seconds=60 * 60 * 12)
            except Exception as e:
                logger.error(f'Failed to set prompt template in Redis cache: {e}')

        return self._cached_template_str

    def render(self, context: Dict[str, Any], force_refresh: bool = False) -> Optional[str]:
        template_str = self.load_template_string(force_refresh=force_refresh)
        if not template_str:
            return None
        try:
            template = self._jinja_env.from_string(template_str)
            return template.render(**context)
        except Exception as e:
            logger.error(f'Failed to render prompt template for "{self.worksheet_name}": {e}')
            return None


class GoogleSheetsPIIPromptStrategy(SpreadsheetPromptStrategy):
    """Backward compatible class for PII detection prompt strategy."""

    def __init__(
        self,
        spreadsheet_url: Optional[str] = None,
        worksheet_name: str = 'PII detection',
        store: Optional[Any] = None,
        cache_key: str = 'pii_detection_prompt_cache',
    ):
        super().__init__(
            worksheet_name=worksheet_name,
            spreadsheet_url=spreadsheet_url,
            store=store,
            cache_key=cache_key,
        )


class GoogleSheetsPIIReflectionPromptStrategy(SpreadsheetPromptStrategy):
    """Backward compatible class for PII reflection prompt strategy."""

    def __init__(
        self,
        spreadsheet_url: Optional[str] = None,
        worksheet_name: str = 'PII reflection',
        store: Optional[Any] = None,
        cache_key: str = 'pii_reflection_prompt_cache',
    ):
        super().__init__(
            worksheet_name=worksheet_name,
            spreadsheet_url=spreadsheet_url,
            store=store,
            cache_key=cache_key,
        )
